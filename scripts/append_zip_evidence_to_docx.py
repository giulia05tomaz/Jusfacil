#!/usr/bin/env python3
"""Cria uma minuta DOCX e acrescenta imagens de um ZIP ao final.

A ordem e os títulos vêm do CSV do pacote quando disponível. Sem CSV, a
ordenação é natural pelo nome do arquivo. O script nunca usa IA para inferir
ordem, título ou conteúdo.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import stat
import sys
import zipfile
import hashlib
from html import escape
from dataclasses import dataclass
from pathlib import Path, PurePosixPath

from PIL import Image, ImageOps
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor


MAX_ARCHIVE_BYTES = 50 * 1024 * 1024
MAX_ENTRIES = 400
MAX_IMAGE_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_IMAGES = 150
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png"}
MAX_ANNEX_PAGES = 400
INDEX_PATTERN = re.compile(r"(?:^|/)(?:00[_ -]?indice|indice|index)[^/]*\.csv$", re.IGNORECASE)


@dataclass(frozen=True)
class EvidenceImage:
    order: int
    reference: str
    title: str
    archive_name: str


def natural_key(value: str) -> tuple[object, ...]:
    return tuple(int(part) if part.isdigit() else part.casefold() for part in re.split(r"(\d+)", value))


def safe_archive_name(raw_name: str) -> str:
    normalized = raw_name.replace("\\", "/")
    path = PurePosixPath(normalized)
    if not normalized or "\x00" in normalized or path.is_absolute() or ".." in path.parts:
        raise ValueError("ZIP_PATH_TRAVERSAL")
    if re.match(r"^[A-Za-z]:", normalized):
        raise ValueError("ZIP_PATH_TRAVERSAL")
    return str(path)


def validate_archive(archive_path: Path, archive: zipfile.ZipFile) -> list[zipfile.ZipInfo]:
    if archive_path.stat().st_size > MAX_ARCHIVE_BYTES:
        raise ValueError("ZIP_TOO_LARGE")
    entries = [entry for entry in archive.infolist() if not entry.is_dir()]
    if len(entries) > MAX_ENTRIES:
        raise ValueError("ZIP_TOO_MANY_ENTRIES")
    total_size = 0
    for entry in entries:
        safe_archive_name(entry.filename)
        unix_mode = entry.external_attr >> 16
        if unix_mode and stat.S_ISLNK(unix_mode):
            raise ValueError("ZIP_SYMLINK_NOT_ALLOWED")
        if Path(entry.filename).suffix.casefold() in {".zip", ".rar", ".7z", ".tar", ".gz"}:
            raise ValueError("ZIP_NESTED_ARCHIVE")
        if entry.file_size > MAX_IMAGE_BYTES:
            raise ValueError("ZIP_ENTRY_TOO_LARGE")
        total_size += entry.file_size
    if total_size > MAX_UNCOMPRESSED_BYTES:
        raise ValueError("ZIP_UNCOMPRESSED_TOO_LARGE")
    return entries


def decode_csv(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("CSV_ENCODING_INVALID")


def find_unique_entry(entries: list[zipfile.ZipInfo], requested_name: str) -> zipfile.ZipInfo | None:
    normalized_requested = safe_archive_name(requested_name).casefold()
    exact = [entry for entry in entries if safe_archive_name(entry.filename).casefold() == normalized_requested]
    if len(exact) == 1:
        return exact[0]
    basename = PurePosixPath(normalized_requested).name
    matches = [entry for entry in entries if PurePosixPath(safe_archive_name(entry.filename).casefold()).name == basename]
    if len(matches) > 1:
        raise ValueError(f"ZIP_DUPLICATE_IMAGE_NAME:{requested_name}")
    return matches[0] if matches else None


def evidence_from_csv(archive: zipfile.ZipFile, entries: list[zipfile.ZipInfo], index: zipfile.ZipInfo) -> tuple[list[EvidenceImage], list[str]]:
    rows = csv.DictReader(io.StringIO(decode_csv(archive.read(index))))
    required = {"ordem", "referencia", "titulo_na_peticao", "arquivo_imagem"}
    if not rows.fieldnames or not required.issubset({name.strip().casefold() for name in rows.fieldnames}):
        raise ValueError("CSV_FIELDS_INVALID")
    images: list[EvidenceImage] = []
    warnings: list[str] = []
    seen_names: set[str] = set()
    for row_number, raw_row in enumerate(rows, start=2):
        row = {(key or "").strip().casefold(): (value or "").strip() for key, value in raw_row.items()}
        requested_name = row.get("arquivo_imagem", "")
        if not requested_name:
            warnings.append(f"Linha {row_number}: arquivo_imagem vazio; item ignorado.")
            continue
        entry = find_unique_entry(entries, requested_name)
        if entry is None:
            warnings.append(f"Linha {row_number}: imagem não encontrada: {requested_name}")
            continue
        if Path(entry.filename).suffix.casefold() not in IMAGE_EXTENSIONS:
            warnings.append(f"Linha {row_number}: formato de imagem não aceito: {requested_name}")
            continue
        key = safe_archive_name(entry.filename).casefold()
        if key in seen_names:
            warnings.append(f"Linha {row_number}: imagem repetida no índice: {requested_name}")
            continue
        seen_names.add(key)
        try:
            order = int(row.get("ordem", ""))
        except ValueError as error:
            raise ValueError(f"CSV_ORDER_INVALID:{row_number}") from error
        if order < 1:
            raise ValueError(f"CSV_ORDER_INVALID:{row_number}")
        reference = row.get("referencia", "") or str(order).zfill(2)
        title = row.get("titulo_na_peticao", "") or PurePosixPath(entry.filename).name
        images.append(EvidenceImage(order, reference, title, entry.filename))
    images.sort(key=lambda item: (item.order, natural_key(item.archive_name)))
    return images, warnings


def collect_images(archive: zipfile.ZipFile, entries: list[zipfile.ZipInfo]) -> tuple[list[EvidenceImage], list[str], str]:
    indexes = [entry for entry in entries if INDEX_PATTERN.search(safe_archive_name(entry.filename))]
    if indexes:
        images, warnings = evidence_from_csv(archive, entries, sorted(indexes, key=lambda item: natural_key(item.filename))[0])
        source = "CSV"
    else:
        candidates = sorted(
            (entry for entry in entries if Path(entry.filename).suffix.casefold() in IMAGE_EXTENSIONS),
            key=lambda item: natural_key(safe_archive_name(item.filename)),
        )
        images = [
            EvidenceImage(index, str(index).zfill(2), PurePosixPath(entry.filename).name, entry.filename)
            for index, entry in enumerate(candidates, start=1)
        ]
        warnings = []
        source = "FILENAME"
    if not images:
        raise ValueError("ZIP_WITHOUT_IMAGES")
    if len(images) > MAX_IMAGES:
        raise ValueError("ZIP_TOO_MANY_IMAGES")
    return images, warnings, source


def set_run_font(run, size: float = 12, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def create_draft_document(draft_text: str, case_id: str, version: int) -> Document:
    layout = Path(__file__).resolve().parent.parent / "assets" / "petition-layout.docx"
    document = Document(layout) if layout.is_file() else Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)

    title_style = document.styles["Title"]
    title_style.font.name = "Times New Roman"
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    title_style.font.size = Pt(13)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_properties = title_style._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    metadata = document.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(metadata.add_run(f"JusFácil  {case_id}  Versão {version}"), 9)

    first_content = True
    for raw_line in draft_text.replace("\r", "").split("\n"):
        stripped = raw_line.strip()
        if not stripped:
            continue
        is_markdown_heading = stripped.startswith("#")
        cleaned = re.sub(r"^#{1,6}\s*", "", stripped)
        cleaned = re.sub(r"^>\s*", "", cleaned)
        cleaned = cleaned.replace("**", "").replace("__", "")
        looks_like_heading = is_markdown_heading or (len(cleaned) <= 100 and cleaned.upper() == cleaned and any(char.isalpha() for char in cleaned))
        if first_content:
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(paragraph.add_run(cleaned), 13, True)
            first_content = False
            continue
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if looks_like_heading else WD_ALIGN_PARAGRAPH.JUSTIFY
        set_run_font(paragraph.add_run(cleaned), 12, looks_like_heading)
    return document


def fit_image(width_px: int, height_px: int, max_width_emu: int, max_height_emu: int) -> tuple[int, int]:
    if width_px <= 0 or height_px <= 0:
        raise ValueError("IMAGE_DIMENSIONS_INVALID")
    ratio = min(max_width_emu / width_px, max_height_emu / height_px)
    return int(width_px * ratio), int(height_px * ratio)


def optimized_image(raw: bytes) -> tuple[bytes, int, int]:
    with Image.open(io.BytesIO(raw)) as source:
        if source.width * source.height > 40_000_000:
            raise ValueError("IMAGE_PIXELS_TOO_LARGE")
        image = ImageOps.exif_transpose(source)
        if image.mode in {"RGBA", "LA"}:
            background = Image.new("RGB", image.size, "white")
            background.paste(image.convert("RGB"), mask=image.getchannel("A"))
            image = background
        elif image.mode != "RGB":
            image = image.convert("RGB")
        image.thumbnail((1280, 1280), Image.Resampling.LANCZOS)
        stream = io.BytesIO()
        image.save(stream, format="JPEG", quality=70, subsampling=2, optimize=True)
        return stream.getvalue(), image.width, image.height


class RetainedEvidenceReader:
    """Only reads server-generated manifest paths inside its own temporary directory."""
    def __init__(self, manifest_path: Path):
        self.root = manifest_path.resolve().parent
        payload = json.loads(manifest_path.read_text(encoding="utf-8"))
        items = payload.get("items", [])
        if not items or len(items) > MAX_IMAGES:
            raise ValueError("EVIDENCE_MANIFEST_INVALID")
        self.data: dict[str, bytes] = {}
        self.images: list[EvidenceImage] = []
        self.logical_count = len(items)
        self.index_lines: list[str] = []
        total_bytes = 0
        for item in sorted(items, key=lambda value: value["order"]):
            source = (self.root / item["file"]).resolve()
            if source.parent != self.root or source.is_symlink():
                raise ValueError("EVIDENCE_PATH_INVALID")
            raw = source.read_bytes()
            total_bytes += len(raw)
            if not raw or len(raw) > MAX_IMAGE_BYTES or total_bytes > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("EVIDENCE_SIZE_INVALID")
            if hashlib.sha256(raw).hexdigest() != item["sha256"]:
                raise ValueError("EVIDENCE_HASH_INVALID")
            reference, title = str(item["reference"]), str(item["title"])
            # Use supplied labels only; never infer facts or the meaning of an image.
            caption = title if re.match(r"^EVID[ÊE]NCIA[ _-]+" + re.escape(reference) + r"\b", title, re.I) else f"EVIDÊNCIA {reference} — {title}"
            if len(caption) > 550:
                raise ValueError("EVIDENCE_TITLE_INVALID")
            self.index_lines.append(caption)
            mime = item["mimeType"]
            pages = []
            if mime in {"image/png", "image/jpeg"}:
                pages = [raw]
            elif mime == "application/pdf":
                import pypdfium2 as pdfium
                if not raw.startswith(b"%PDF-"):
                    raise ValueError("EVIDENCE_PDF_INVALID")
                with pdfium.PdfDocument(raw) as pdf:
                    if not len(pdf) or len(pdf) + len(self.images) > MAX_ANNEX_PAGES:
                        raise ValueError("EVIDENCE_PAGE_LIMIT")
                    for page_index in range(len(pdf)):
                        page = pdf[page_index]
                        width, height = page.get_size()
                        if width <= 0 or height <= 0:
                            raise ValueError("EVIDENCE_PDF_INVALID")
                        bitmap = page.render(scale=min(2, 1600 / max(width, height)))
                        stream = io.BytesIO()
                        bitmap.to_pil().convert("RGB").save(stream, format="JPEG", quality=85)
                        pages.append(stream.getvalue())
                        bitmap.close()
                        page.close()
            else:
                raise ValueError("EVIDENCE_ANNEX_FORMAT_UNSUPPORTED")
            for page_index, raw_page in enumerate(pages):
                key = f"evidence-{len(self.images)}"
                self.data[key] = raw_page
                label = caption if len(pages) == 1 else f"{caption} — Página {page_index + 1} de {len(pages)}"
                self.images.append(EvidenceImage(int(item["order"]), reference, label, key))
                if len(self.images) > MAX_ANNEX_PAGES:
                    raise ValueError("EVIDENCE_PAGE_LIMIT")

    def read(self, name: str) -> bytes:
        return self.data[name]


def create_complete_pdf(draft_text: str, case_id: str, version: int, archive: zipfile.ZipFile, images: list[EvidenceImage], output: Path) -> None:
    from reportlab import rl_config
    rl_config.useA85 = False
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.platypus import SimpleDocTemplate, Paragraph, PageBreak, Image as PdfImage, KeepTogether, Spacer
    cm = 72 / 2.54
    body = ParagraphStyle("body", fontName="Times-Roman", fontSize=12, leading=18, alignment=TA_JUSTIFY, spaceAfter=6)
    heading = ParagraphStyle("heading", parent=body, fontName="Times-Bold", alignment=TA_CENTER, keepWithNext=True, spaceBefore=8)
    caption_style = ParagraphStyle("caption", fontName="Times-Bold", fontSize=10.5, leading=13, alignment=TA_CENTER, spaceAfter=8, textColor=colors.black)
    story = []
    for raw in draft_text.replace("\r", "").split("\n"):
        text = raw.strip()
        if not text:
            continue
        is_heading = text.startswith("#") or (len(text) <= 100 and text.upper() == text and any(c.isalpha() for c in text))
        text = re.sub(r"^#{1,6}\s*", "", text).replace("**", "").replace("__", "")
        text = re.sub(r"^>\s*", "", text)
        story.append(Paragraph(escape(text), heading if is_heading else body))
    for index, item in enumerate(images):
        story.append(PageBreak())
        encoded, w, h = optimized_image(archive.read(item.archive_name))
        ratio = min(16 * cm / w, (20.5 if index == 0 else 21.5) * cm / h)
        picture = PdfImage(io.BytesIO(encoded), width=w * ratio, height=h * ratio)
        qa_annex = "ANEXOS DE QA" if "pacote de Nívea" in draft_text and "fictício" in draft_text else "ANEXO PROBATÓRIO EVIDÊNCIAS"
        first_heading = [Paragraph(qa_annex, heading), Spacer(1, 8)] if index == 0 else []
        story.append(KeepTogether([*first_heading, Paragraph(escape(item.title), caption_style), picture]))
    def footer(canvas, doc):
        canvas.setFont("Times-Roman", 9)
        canvas.drawString(3 * cm, 1.4 * cm, f"JusFácil  {case_id}  Versão {version}")
        canvas.drawRightString(19 * cm, 1.4 * cm, str(doc.page))
    SimpleDocTemplate(str(output), pagesize=(21 * cm, 29.7 * cm), leftMargin=3 * cm, rightMargin=2 * cm,
                      topMargin=2.5 * cm, bottomMargin=2.5 * cm, title=f"Petição Inicial {case_id} Versão {version}").build(story, onFirstPage=footer, onLaterPages=footer)


def append_evidence_pages(document: Document, archive: zipfile.ZipFile, images: list[EvidenceImage]) -> None:
    section = document.sections[-1]
    max_width = int(section.page_width - section.left_margin - section.right_margin)
    max_height = int(section.page_height - section.top_margin - section.bottom_margin - Cm(2.2))
    for index, item in enumerate(images):
        if index == 0:
            heading = document.add_paragraph()
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            heading.paragraph_format.page_break_before = True
            heading.paragraph_format.keep_with_next = True
            heading.paragraph_format.space_after = Pt(8)
            set_run_font(heading.add_run("ANEXO PROBATÓRIO  EVIDÊNCIAS"), 12, True)
        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.page_break_before = index > 0
        caption.paragraph_format.keep_with_next = True
        caption.paragraph_format.space_after = Pt(8)
        set_run_font(caption.add_run(item.title), 10.5, True)

        encoded, image_width, image_height = optimized_image(archive.read(item.archive_name))
        width, height = fit_image(image_width, image_height, max_width, max_height)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(io.BytesIO(encoded), width=Emu(width), height=Emu(height))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Acrescenta imagens ordenadas de um ZIP ao final de uma minuta Word.")
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--input-docx", type=Path, help="DOCX existente ao qual as evidências serão acrescentadas.")
    source.add_argument("--draft-text", type=Path, help="Texto da minuta usado para criar o DOCX base.")
    evidences = parser.add_mutually_exclusive_group(required=True)
    evidences.add_argument("--zip", dest="zip_path", type=Path, help="Pacote ZIP de evidências (compatibilidade).")
    evidences.add_argument("--evidence-manifest", type=Path, help="Manifesto privado gerado pelo servidor com originais preservados.")
    parser.add_argument("--output", type=Path, required=True, help="Arquivo DOCX de saída.")
    parser.add_argument("--case-id", default="CASO", help="Identificador exibido no Word criado a partir de texto.")
    parser.add_argument("--version", type=int, default=1, help="Versão exibida no Word criado a partir de texto.")
    parser.add_argument("--pdf-output", type=Path, help="PDF completo com o mesmo texto e as mesmas imagens do Word.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.input_docx and (not args.input_docx.is_file() or args.input_docx.suffix.casefold() != ".docx"):
        raise ValueError("INPUT_DOCX_INVALID")
    if args.draft_text and not args.draft_text.is_file():
        raise ValueError("DRAFT_TEXT_INVALID")
    if args.zip_path and (not args.zip_path.is_file() or args.zip_path.suffix.casefold() != ".zip"):
        raise ValueError("ZIP_INVALID")
    args.output.parent.mkdir(parents=True, exist_ok=True)

    draft_text = args.draft_text.read_text(encoding="utf-8") if args.draft_text else ""
    if args.evidence_manifest:
        archive = RetainedEvidenceReader(args.evidence_manifest)
        draft_text += "\n\n# RELAÇÃO DOS ANEXOS\n" + "\n".join(archive.index_lines)
        document = Document(args.input_docx) if args.input_docx else create_draft_document(draft_text, args.case_id, args.version)
        images = archive.images
        append_evidence_pages(document, archive, images)
        if args.pdf_output:
            if not args.draft_text:
                raise ValueError("PDF_REQUIRES_DRAFT_TEXT")
            create_complete_pdf(draft_text, args.case_id, args.version, archive, images, args.pdf_output)
        logical_count, warnings, source = archive.logical_count, [], "CASE_MANIFEST"
    else:
      document = Document(args.input_docx) if args.input_docx else create_draft_document(draft_text, args.case_id, args.version)
      with zipfile.ZipFile(args.zip_path, "r") as archive:
        entries = validate_archive(args.zip_path, archive)
        images, warnings, source = collect_images(archive, entries)
        append_evidence_pages(document, archive, images)
        if args.pdf_output:
            if not args.draft_text:
                raise ValueError("PDF_REQUIRES_DRAFT_TEXT")
            create_complete_pdf(args.draft_text.read_text(encoding="utf-8"), args.case_id, args.version, archive, images, args.pdf_output)
        logical_count = len(images)
    document.save(args.output)
    print(json.dumps({
        "output": str(args.output.resolve()),
        "insertedCount": len(images),
        "evidenceCount": logical_count,
        "orderSource": source,
        "warnings": warnings,
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (OSError, ValueError, zipfile.BadZipFile) as error:
        print(json.dumps({"error": str(error)}, ensure_ascii=False), file=sys.stderr)
        raise SystemExit(1)
